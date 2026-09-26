from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

from fpdf import FPDF


def sanitize_text(text: str) -> str:
    """
    Cleans text before exporting.
    """

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.strip()


def format_txt(text: str) -> bytes:
    """
    Returns UTF-8 TXT content.
    """

    clean_text = sanitize_text(text)

    return clean_text.encode("utf-8")


def format_docx(
    text: str,
    doc_type: str = "Legal Document"
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    # Default font
    styles = document.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)

    # Title
    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        sanitize_text(doc_type).upper()
    )

    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    # Body
    clean_text = sanitize_text(text)

    for line in clean_text.splitlines():

        line = line.strip()

        if not line:
            document.add_paragraph("")
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

        run = paragraph.add_run(line)

        run.font.name = "Times New Roman"
        run.font.size = Pt(12)

        upper = line.upper()

        if (
            upper in {
                "PARTIES",
                "AGREEMENT",
                "TERMS AND CONDITIONS",
                "GENERAL PROVISIONS",
                "SIGNATURES",
            }
            or (
                len(line) < 100
                and line[:2].isdigit()
                and "." in line[:4]
            )
        ):
            run.bold = True

    # Footer
    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_paragraph.add_run(
        "LegalEase - AI-assisted legal document draft"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(9)

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.read()


class LegalEasePDF(FPDF):

    def __init__(self, doc_type: str):
        super().__init__()

        self.doc_type = doc_type

        self.set_auto_page_break(
            auto=True,
            margin=20
        )

    def header(self):

        self.set_font(
            "Times",
            "B",
            12
        )

        self.cell(
            0,
            10,
            "LegalEase",
            align="C"
        )

        self.ln(8)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Times",
            "",
            9
        )

        self.cell(
            0,
            10,
            f"LegalEase | Page {self.page_no()}",
            align="C"
        )


def format_pdf(
    text: str,
    doc_type: str = "Legal Document"
) -> bytes:

    pdf = LegalEasePDF(doc_type)

    pdf.set_title(doc_type)

    pdf.add_page()

    # Document title
    pdf.set_font(
        "Times",
        "B",
        16
    )

    pdf.multi_cell(
        0,
        10,
        sanitize_text(doc_type).upper(),
        align="C"
    )

    pdf.ln(5)

    # Document content
    pdf.set_font(
        "Times",
        "",
        12
    )

    clean_text = sanitize_text(text)

    for line in clean_text.splitlines():

        line = line.strip()

        if not line:
            pdf.ln(5)
            continue

        is_heading = (
            line.upper()
            in {
                "PARTIES",
                "AGREEMENT",
                "TERMS AND CONDITIONS",
                "GENERAL PROVISIONS",
                "SIGNATURES",
            }
        )

        if is_heading:
            pdf.set_font(
                "Times",
                "B",
                12
            )

        else:
            pdf.set_font(
                "Times",
                "",
                12
            )

        pdf.multi_cell(
            0,
            7,
            line
        )

        pdf.ln(2)

    pdf.ln(5)

    pdf.set_font(
        "Times",
        "I",
        9
    )

    pdf.multi_cell(
        0,
        5,
        (
            "Notice: This document is an AI-assisted draft "
            "and should be reviewed by a qualified legal "
            "professional before use."
        )
    )

    output = pdf.output()

    return bytes(output)


def format_html_preview(text: str) -> str:
    """
    Converts plain text into safe HTML preview content.
    """

    clean_text = sanitize_text(text)

    html_parts = []

    for line in clean_text.splitlines():

        if not line.strip():
            html_parts.append("<br>")
            continue

        safe_line = escape(line)

        upper = line.upper()

        if upper in {
            "PARTIES",
            "AGREEMENT",
            "TERMS AND CONDITIONS",
            "GENERAL PROVISIONS",
            "SIGNATURES",
        }:
            html_parts.append(
                f"<h3>{safe_line}</h3>"
            )

        else:
            html_parts.append(
                f"<p>{safe_line}</p>"
            )

    return "\n".join(html_parts)